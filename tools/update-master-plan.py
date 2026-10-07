from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "outputs" / "master-plan" / "Master Plan.xlsx"
DONE = "Đã hoàn thành"

completed_plan_ids = {
    "2.2", "2.3", "2.4", "2.5", "2.6", "2.7",
    "3.1", "3.2", "3.3", "3.4", "3.5",
    "4.1", "4.2", "4.3", "4.4", "4.5", "4.6", "4.7", "4.8",
    "5.1", "5.2", "5.3",
    "6.1", "6.2", "6.3",
    "7.1", "7.2", "7.3",
    "8.1", "8.2", "8.3", "8.4", "8.5",
    "9.1", "9.2", "9.3",
}

notes = {
    "2.2": "Menu API CRUD, validation, soft-delete and owner RBAC.",
    "2.3": "Variants, topping groups, topping links and min/max validation.",
    "2.4": "Catalog search/filter/sort/pagination and customer search history.",
    "2.5": "Promotion, banner and favorite flows.",
    "2.6": "Realistic seed with four roles, restaurants, menu items and operational data.",
    "2.7": "JPEG/PNG/WebP image upload and square crop/rotate/zoom/reset editor.",
    "3.1": "Customer/restaurant cart with quantity, topping, note and price snapshot.",
    "3.2": "Transactional checkout with server pricing and idempotency key.",
    "3.3": "Restaurant order state machine with status history and rejection.",
    "3.4": "Customer ordering UI, address/payment/promotion and notifications.",
    "3.5": "Checkout-confirm-prepare-ready workflow, RBAC, audit and regression tests.",
    "4.1": "Shipper profile, availability, assignment offer/accept/reject and permissions.",
    "4.2": "Delivery lifecycle, status history and delivery zones.",
    "4.3": "GPS point ingestion and delivery tracking API.",
    "4.4": "Payment/COD flow, callback and failure handling.",
    "4.5": "Invoice creation and order snapshot lookup.",
    "4.6": "Notifications, order chat, support ticket and reviews.",
    "4.7": "Cancellation, refund and customer wallet.",
    "4.8": "Shipper earnings/withdrawals and restaurant settlement/commission.",
    "5.1": "CSV import for RESTAURANTS/MENU_ITEMS with import_jobs and per-row results; smoke import passed.",
    "5.2": "ORDERS/REVENUE/SHIPPER/SETTLEMENT export as CSV, Excel SpreadsheetML and PDF; export_jobs and /exports.",
    "5.3": "background_jobs worker, 30-second polling, three-attempt retry and token cleanup; run endpoint passed.",
    "6.1": "Helmet, allow-list CORS, body limits, rate limit, validation, upload whitelist, RBAC and audit.",
    "6.2": "Catalog TTL response cache, hit/miss header and menu mutation invalidation; Redis is included in Compose.",
    "6.3": "Structured JSON HTTP logs, readiness/liveness, audit middleware and system_settings admin API.",
    "7.1": "63 unit/service tests cover auth, address, catalog, menu, order state and RBAC.",
    "7.2": "API plus MySQL smoke covers login/me, cache, import/export, worker, transaction/idempotency and backup.",
    "7.3": "Realistic seed contains 20 restaurants and 4,200 menu items; Node coverage is 37.04%.",
    "8.1": "Docker Compose API, MySQL 8, Redis 7 and MailHog with environment, init SQL and healthcheck.",
    "8.2": "GitHub Actions check/test workflow, production Dockerfile and package:release script.",
    "8.3": "OpenAPI 3.0.3 at /api/openapi.json and tested Postman collection for all four roles.",
    "8.4": "Deployment guide, user guide, handover checklist, backend README and OpenAPI contract.",
    "8.5": "Smoke script, realistic seed, sample data, export/import artifacts and generated release package.",
    "9.1": "Smoke after API restart, readiness/catalog/cache checks, GPS/payment route inventory and verified mysqldump.",
    "9.2": "Fixed background enqueue, token cleanup column, Docker storage path and regressions; 63 tests pass.",
    "9.3": "Release package includes backend/frontend/SQL/docs/Postman plus handover checklist and test evidence.",
}

workbook = load_workbook(WORKBOOK)
plan = workbook["Plan"]
for row in range(1, plan.max_row + 1):
    task_id = str(plan.cell(row, 1).value or "")
    if task_id in completed_plan_ids:
        plan.cell(row, 6).value = DONE
        plan.cell(row, 7).value = 1

coverage = workbook["SQL coverage"]
for row in range(1, coverage.max_row + 1):
    plan_id = str(coverage.cell(row, 4).value or "")
    ids = {value.strip() for value in plan_id.split(",")}
    matching = sorted(ids & completed_plan_ids)
    if matching:
        coverage.cell(row, 5).value = DONE
        coverage.cell(row, 6).value = "; ".join(notes[item] for item in matching if item in notes)

summary = workbook["Summary"]
summary.cell(11, 2).value = (
    "Core 2.2-4.8 and operations 5.1-9.3 completed: import/export, background jobs, "
    "security/cache/logging, tests, Docker/CI, OpenAPI/docs, smoke/backup/release."
)
summary.cell(12, 2).value = (
    "Redis and MailHog are optional Docker services; local development keeps an in-memory "
    "cache fallback so the single-port API runs without Redis."
)
summary.cell(5, 9).value = "9.3/20"
summary.cell(2, 1).value = (
    "Implementation plan based on food_delivery_db.sql and assignment requirements | Updated: 07/10/2026"
)

infrastructure = workbook["Ha tang"]
infrastructure.cell(6, 4).value = (
    "Four-role web/API plus import/export/jobs, security/cache/logging, Docker/CI, "
    "OpenAPI/Postman, smoke/backup/release are available; Redis/MailHog are in Compose."
)

workbook.calculation.fullCalcOnLoad = True
workbook.calculation.forceFullCalc = True
workbook.calculation.calcMode = "auto"
workbook.save(WORKBOOK)
print(f"Updated {WORKBOOK} at {datetime.now().isoformat(timespec='seconds')}")
